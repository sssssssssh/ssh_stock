# ruff: noqa: E501
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]


PRD_PROGRESS = [
    ("23. 实现进度（2026-08-26）", 1),
    ("23.1 当前已落地范围", 2),
    (
        "本轮按系统设计中的 Codex 首条开发指令，仅完成 Milestone 0 和 Milestone 1，"
        "不提前实现页面、策略评分和状态机。",
        0,
    ),
    ("已创建后端工程骨架：Python 3.12、FastAPI、SQLAlchemy 2.x、Alembic、pytest。", "bullet"),
    ("已创建 Docker Compose：PostgreSQL 16、backend、worker。", "bullet"),
    ("已实现配置系统：.env + config/app.yaml + config/strategy.yaml。", "bullet"),
    ("已实现 MarketDataProvider 抽象与 TushareProvider。", "bullet"),
    (
        "已实现 stock_basic、trade_calendar、stock_daily、stock_adj_factor、"
        "stock_daily_basic、index_daily 原始数据表。",
        "bullet",
    ),
    ("已实现 job_run 和 provider_api_log 运行日志表。", "bullet"),
    ("已实现 daily / backfill / scheduler CLI。", "bullet"),
    ("已实现原始日线数据的基础质量检查。", "bullet"),
    ("已实现 pytest 单元测试与 API health 导入测试。", "bullet"),
    ("23.2 验证结果", 2),
    ("python -m pytest：9 passed，1 个 FastAPI/Starlette TestClient 第三方弃用警告。", "bullet"),
    ("python -m ruff check .：All checks passed。", "bullet"),
    ("python -m alembic upgrade head --sql：离线 SQL 生成成功。", "bullet"),
    ("python -m app.cli --help：CLI 入口正常。", "bullet"),
    ("23.3 尚未实现", 2),
    ("Milestone 2：复权价格、基础因子、Eligible Universe、RPS。", "bullet"),
    ("Milestone 3：Market Score、Sector Heat、Lifecycle。", "bullet"),
    ("Milestone 4：RightSideScore、TrendScore、S0-S6 状态机、信号生成。", "bullet"),
    ("Milestone 5：完整 REST API。", "bullet"),
    ("Milestone 6：Vue 前端。", "bullet"),
    ("Milestone 7：后验统计和研究接口。", "bullet"),
    ("23.4 换电脑继续开发要求", 2),
    (
        "新电脑继续开发前，先阅读 README.md 和 PROJECT_RULES.md；后续 Codex 开发必须"
        "先完成当前 Milestone 的验收标准，再进入下一 Milestone。",
        0,
    ),
]


DESIGN_PROGRESS = [
    ("35. 实现进度（2026-08-26）", 1),
    ("35.1 当前代码结构", 2),
    (
        "本轮已在当前目录初始化项目工程，完成 Milestone 0 和 Milestone 1 的首版可运行底座。",
        0,
    ),
    ("README.md、PROJECT_RULES.md、pyproject.toml、docker-compose.yml、Dockerfile、alembic.ini", "bullet"),
    ("config/app.yaml、config/strategy.yaml", "bullet"),
    ("migrations/env.py、migrations/versions/20260826_0001_initial_raw_warehouse.py", "bullet"),
    ("backend/app/main.py、cli.py、core、api/v1、models、providers、repositories、services、jobs", "bullet"),
    ("tests/unit、tests/integration", "bullet"),
    ("35.2 已实现模块", 2),
    ("app.main：FastAPI 应用工厂与 /health。", "bullet"),
    ("app.api.v1.system：GET /api/v1/system/status，返回核心原始表最新日期。", "bullet"),
    ("app.api.v1.jobs：GET /api/v1/jobs，返回最近任务运行记录。", "bullet"),
    ("app.core.config：Pydantic Settings + YAML 配置合并。", "bullet"),
    ("app.models.market_data：P0 原始行情表。", "bullet"),
    ("app.models.job：任务日志与 Provider API 日志。", "bullet"),
    ("app.providers.base：数据源协议。", "bullet"),
    ("app.providers.tushare_provider：Tushare SDK 适配，日期在此转换为 YYYYMMDD。", "bullet"),
    ("app.providers.logging_provider：Provider 调用入库日志包装器。", "bullet"),
    ("app.services.ingestion：标准化、质量检查、upsert 入库。", "bullet"),
    ("app.jobs.daily_job / backfill_job / scheduler：任务编排。", "bullet"),
    ("app.cli：daily / backfill / scheduler 命令，后续 Milestone 命令以明确错误提示保留。", "bullet"),
    ("35.3 迁移与表", 2),
    ("首个 Alembic revision：0001_initial_raw_warehouse。", 0),
    (
        "已建表：stock_basic、trade_calendar、stock_daily、stock_adj_factor、"
        "stock_daily_basic、index_daily、job_run、provider_api_log。",
        "bullet",
    ),
    ("所有原始事实表通过 PostgreSQL INSERT ... ON CONFLICT DO UPDATE 实现幂等写入。", "bullet"),
    ("35.4 验证结果", 2),
    ("python -m pytest：9 passed，1 个第三方弃用警告。", "bullet"),
    ("python -m ruff check .：All checks passed。", "bullet"),
    ("python -m alembic upgrade head --sql：成功生成 PostgreSQL DDL。", "bullet"),
    ("python -m app.cli --help：CLI 入口正常。", "bullet"),
    ("35.5 下一步开发入口", 2),
    ("下一轮应进入 Milestone 2：Factor Engine。", 0),
    ("读取最近 300~320 个交易日 OHLCV + adj_factor。", "bullet"),
    ("计算 adjusted OHLC。", "bullet"),
    ("实现 MA、return、slope、ATR、breakout、higher-low、drawdown、trend efficiency。", "bullet"),
    ("实现 Eligible Universe 和 RPS。", "bullet"),
    ("新增 stock_factor_daily 模型、迁移、upsert 和测试。", "bullet"),
]


def append_progress(path: Path, content: list[tuple[str, int | str]]) -> None:
    doc = Document(path)
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
    existing_text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    marker = content[0][0]
    if marker in existing_text:
        doc.save(path)
        return

    doc.add_page_break()
    for text, style in content:
        paragraph = doc.add_paragraph()
        if style == 1:
            run = paragraph.add_run(text)
            run.bold = True
            run.font.size = Pt(16)
        elif style == 2:
            run = paragraph.add_run(text)
            run.bold = True
            run.font.size = Pt(13)
        elif style == "bullet":
            paragraph.paragraph_format.left_indent = Pt(18)
            paragraph.paragraph_format.first_line_indent = Pt(-9)
            run = paragraph.add_run("- " + text)
            run.font.size = Pt(10.5)
        else:
            run = paragraph.add_run(text)
            run.font.size = Pt(10.5)
    doc.save(path)


def append_dependency_note(path: Path) -> None:
    doc = Document(path)
    existing_text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    marker = "依赖安装补充（2026-08-26）"
    if marker in existing_text:
        return

    paragraph = doc.add_paragraph()
    run = paragraph.add_run(marker)
    run.bold = True
    run.font.size = Pt(13)

    paragraph = doc.add_paragraph()
    run = paragraph.add_run(
        "项目依赖以 pyproject.toml 为主，同时补充 requirements.txt 和 "
        "requirements-dev.txt，便于在新电脑上使用传统 pip 命令安装。"
    )
    run.font.size = Pt(10.5)
    doc.save(path)


def append_tushare_diagnostic_note(path: Path) -> None:
    doc = Document(path)
    existing_text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    marker = "Tushare 配置诊断补充（2026-08-27）"
    if marker in existing_text:
        return

    paragraph = doc.add_paragraph()
    run = paragraph.add_run(marker)
    run.bold = True
    run.font.size = Pt(13)

    paragraph = doc.add_paragraph()
    run = paragraph.add_run(
        "CLI 已补充 check-config 和 check-tushare。check-config 用于脱敏确认 .env、"
        "DATABASE_URL 和 TUSHARE_TOKEN 是否被读取；check-tushare 使用 trade_cal 小请求"
        "验证 token 是否被 Tushare 服务端接受。"
    )
    run.font.size = Pt(10.5)
    doc.save(path)


def append_milestone2_note(path: Path) -> None:
    doc = Document(path)
    existing_text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    marker = "Milestone 2 实现补充（2026-08-27）"
    if marker in existing_text:
        return

    paragraph = doc.add_paragraph()
    run = paragraph.add_run(marker)
    run.bold = True
    run.font.size = Pt(13)

    items = [
        "已新增 stock_factor_daily 因子表和 Alembic 迁移 0002_stock_factor_daily。",
        "已实现 Factor Engine：复权 OHLC、MA、收益率、斜率、ATR、突破、higher-low、回撤、趋势效率、Eligible Universe、RPS。",
        "已实现 recalc-factors CLI，并将因子重算接入 daily / backfill 任务链路。",
        "系统状态接口已增加 latest_factor_date，便于确认因子表最新日期。",
        "当前验证结果为 python -m pytest：12 passed；python -m ruff check .：All checks passed。",
        "下一步进入 Milestone 3：Market + Sector。",
    ]
    for item in items:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Pt(18)
        paragraph.paragraph_format.first_line_indent = Pt(-9)
        run = paragraph.add_run("- " + item)
        run.font.size = Pt(10.5)
    doc.save(path)


def append_milestone3_note(path: Path) -> None:
    doc = Document(path)
    existing_text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    marker = "Milestone 3 实现补充（2026-08-27）"
    if marker in existing_text:
        return

    paragraph = doc.add_paragraph()
    run = paragraph.add_run(marker)
    run.bold = True
    run.font.size = Pt(13)

    items = [
        "已新增 market_daily、sector、sector_member、sector_factor_daily 表和 Alembic 迁移 0003_market_sector。",
        "已新增 stock_factor_daily.return1 / return3 和 Alembic 迁移 0004_stock_factor_short_returns，用于支撑行业短周期收益聚合。",
        "已实现 Market Score V1：指数趋势、市场宽度、涨跌家数、新高新低、成交额活跃度，并输出 Regime。",
        "已实现申万行业元数据与行业成分同步，行业成员写入通过 PostgreSQL upsert 保持幂等。",
        "已实现 Sector Heat V1：行业收益、超额收益、宽度、RPS、成交活跃度、Heat Momentum、Heat Rank、Lifecycle。",
        "已实现 recalc-market / recalc-sectors CLI，并将市场温度和行业热度计算接入 daily / backfill 任务链路。",
        "系统状态接口已新增 latest_market_date 和 latest_sector_factor_date。",
        "当前验证结果为 python -m pytest：17 passed；python -m ruff check backend tests：All checks passed；python -m alembic heads：0004_stock_factor_short_returns。",
        "下一步进入 Milestone 4：RightSideScore、TrendScore、S0-S6 状态机与信号生成。",
    ]
    for item in items:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Pt(18)
        paragraph.paragraph_format.first_line_indent = Pt(-9)
        run = paragraph.add_run("- " + item)
        run.font.size = Pt(10.5)
    doc.save(path)


def append_tushare_proxy_note(path: Path) -> None:
    doc = Document(path)
    existing_text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    marker = "Tushare 代理初始化修正（2026-08-27）"
    if marker in existing_text:
        return

    paragraph = doc.add_paragraph()
    run = paragraph.add_run(marker)
    run.bold = True
    run.font.size = Pt(13)

    items = [
        "TushareProvider 已调整为代理版 SDK 初始化方式：先 ts.set_token(TUSHARE_TOKEN)，再无参数 ts.pro_api()，最后设置 _DataApi__http_url。",
        "新增配置项 TUSHARE_HTTP_URL，默认代理地址为 https://fastapic.stockai888.top；check-config 会显示当前加载的代理地址。",
        "Tushare SDK 仍只允许出现在 backend/app/providers/tushare_provider.py，业务服务继续通过 MarketDataProvider 抽象调用。",
        "本机已使用 python -m app.cli check-tushare 验证代理连接成功，返回 tushare_ok=true rows=10；全量测试结果为 python -m pytest：17 passed。",
    ]
    for item in items:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Pt(18)
        paragraph.paragraph_format.first_line_indent = Pt(-9)
        run = paragraph.add_run("- " + item)
        run.font.size = Pt(10.5)
    doc.save(path)


def append_daily_fix_note(path: Path) -> None:
    doc = Document(path)
    existing_text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    marker = "Daily 执行排障修复（2026-08-27）"
    if marker in existing_text:
        return

    paragraph = doc.add_paragraph()
    run = paragraph.add_run(marker)
    run.bold = True
    run.font.size = Pt(13)

    items = [
        "已将云端 PostgreSQL 从 0002_stock_factor_daily 升级到 0004_stock_factor_short_returns，补齐市场、行业表和 return1/return3 字段。",
        "已修复 PostgreSQL 单条 SQL 参数数量上限问题，upsert_rows 会按参数数量自动分批写入。",
        "已兼容代理接口申万行业口径：index_classify 使用 SW2021，index_member_all 使用 SW 返回的 l1_code 映射一级行业。",
        "python -m app.cli daily --trade-date 2026-08-26 已执行成功，写入 stock_daily 5547 行、stock_daily_basic 5547 行、index_daily 3 行、stock_factor_daily 5547 行、market_daily 1 行。",
        "当前仅有单日行情，Eligible Universe 为 0，所以 sector_factor_daily 暂为 0 行；完成历史回填后再重算行业热度。",
        "当前验证结果为 python -m pytest：22 passed；python -m ruff check backend tests scripts/sync_docx_progress.py：All checks passed。",
    ]
    for item in items:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Pt(18)
        paragraph.paragraph_format.first_line_indent = Pt(-9)
        run = paragraph.add_run("- " + item)
        run.font.size = Pt(10.5)
    doc.save(path)


def append_milestone4_note(path: Path) -> None:
    doc = Document(path)
    existing_text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    marker = "Milestone 4 实现补充（2026-08-27）"
    if marker in existing_text:
        return

    paragraph = doc.add_paragraph()
    run = paragraph.add_run(marker)
    run.bold = True
    run.font.size = Pt(13)

    items = [
        "已新增 stock_state_daily 状态表和 strategy_signal 信号表，Alembic 当前单头为 0005_trend_state_signal。",
        "已实现 RightSideScore V1：Breakout、MA20 Turn、MA Recovery、RPS Acceleration、Volume Confirmation、Higher Low、Volatility Structure。",
        "已实现 TrendScore V1：RPS20、RPS60、RPS120、MA Structure、Slope、Trend Efficiency、Drawdown Quality。",
        "已实现 S0-S6 状态机、转移矩阵、S3 最长保留天数、S0 快速跳转限制和 OpportunityScore。",
        "已实现 RIGHT_SIDE_NEW、TREND_ENTER、MAIN_UP_ENTER、TREND_DECAY、LEADER_BREAKOUT 信号生成。",
        "recalc-states 已从占位命令改为可执行命令，并接入 daily / backfill 任务链路。",
        "系统状态接口已新增 latest_state_date 和 latest_signal_date。",
        "云端数据库已执行 python -m alembic upgrade head，当前版本为 0005_trend_state_signal。",
        "已使用 2026-08-26 单日数据验证 recalc-states，写入 stock_state_daily 5547 行，strategy_signal 0 行；由于仅有单日历史，当前全部为基础 S0 状态。",
        "当前验证结果为 python -m pytest：25 passed；python -m ruff check backend tests scripts/sync_docx_progress.py：All checks passed。",
        "下一步进入 Milestone 5：完整 REST API。",
    ]
    for item in items:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Pt(18)
        paragraph.paragraph_format.first_line_indent = Pt(-9)
        run = paragraph.add_run("- " + item)
        run.font.size = Pt(10.5)
    doc.save(path)


def main() -> None:
    prd = ROOT / "docs" / "股票机会发现系统_PRD_V1.0.docx"
    design = ROOT / "docs" / "股票机会发现系统_系统设计_V1.0.docx"
    append_progress(prd, PRD_PROGRESS)
    append_progress(design, DESIGN_PROGRESS)
    append_dependency_note(prd)
    append_dependency_note(design)
    append_tushare_diagnostic_note(prd)
    append_tushare_diagnostic_note(design)
    append_milestone2_note(prd)
    append_milestone2_note(design)
    append_milestone3_note(prd)
    append_milestone3_note(design)
    append_tushare_proxy_note(prd)
    append_tushare_proxy_note(design)
    append_daily_fix_note(prd)
    append_daily_fix_note(design)
    append_milestone4_note(prd)
    append_milestone4_note(design)


if __name__ == "__main__":
    main()
